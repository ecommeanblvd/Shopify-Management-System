# -*- coding: utf-8 -*-
"""Sửa cột ngày trong hai file đối soát theo nguồn ngày đã có nhân chứng.

Vì sao (19/10/2026): bản dựng file đầu tiên đọc ngày sai ở HAI chỗ độc lập.
  1. Ô ngày của bảng ship hộ Lark là epoch, mọi dòng TINH kỳ cũ có epoch đúng 17:00:00Z —
     nửa đêm GIỜ VIỆT NAM. Quy về UTC nên LÙI MỘT NGÀY ở mọi dòng.
  2. Sáu ô "Ngày Request" trên sheet WholeSale là SERIAL, và bản đọc đầu biến chúng thành
     chuỗi vô nghĩa kiểu "2026-07-04 (serial)".

Nguồn ngày:
  ngay-lark-<brand>.json      — bảng ship hộ Lark, đọc qua `ngayISO` (+07). Dùng cho đơn bán lẻ.
  wholesale-lark.json         — sheet WholeSale (cả hai brand), serial đã quy đổi + đảo
                                ngày↔tháng có nhân chứng. Dùng cho đơn WholeSale.

Hai nguồn trùng nhau ở 5 mã TINH và khớp cả 5 — script DỪNG nếu có mã nào lệch.

  python3 scripts/du-lieu/sua-ngay-file-doi-soat.py
"""
import json
import re
import openpyxl

MA_DON = re.compile(r'^\d{2}-INSLG-SV-\d{4}$')

# file, sheet, hàng tiêu đề, cột ngày (1-based), nguồn ngày ưu tiên
VIEC = [
    ('doi-soat-tinh-atelier.xlsx', 'Chi tiết 26 đơn', 4, 5, 'shiphho'),
    ('doi-soat-tinh-atelier.xlsx', 'WholeSale TINH', 5, 2, 'wholesale'),
    ('doi-soat-mirer-wholesale.xlsx', 'WholeSale Mirer', 5, 2, 'wholesale'),
]


def nap(p):
    with open(p, encoding='utf-8') as f:
        return json.load(f)


def main():
    ship = {}
    for b in ('tinh', 'mirer'):
        ship.update(nap(f'scripts/du-lieu/ngay-lark-{b}.json'))
    ws_raw = nap('scripts/du-lieu/wholesale-lark.json')['brand']
    ws_don = [d for ds in ws_raw.values() for d in ds]
    ws_ngay = {d['ma']: d['ngay'] for d in ws_don}
    suy_luan = {d['ma'] for d in ws_don if d.get('suyLuanNgay')}

    lech = [f'{k}: ship hộ {ship[k]} ≠ sheet WholeSale {ws_ngay[k]}'
            for k in set(ship) & set(ws_ngay) if ship[k] != ws_ngay[k]]
    if lech:
        raise SystemExit('DỪNG: hai nguồn ngày lệch nhau:\n  ' + '\n  '.join(lech))
    print(f'Hai nguồn trùng {len(set(ship) & set(ws_ngay))} mã, khớp hết.\n')

    for path, ten_sheet, hang_dau, cot, uu_tien in VIEC:
        wb = openpyxl.load_workbook(path)
        sh = wb[ten_sheet]
        doi, thieu = [], []
        for r in range(hang_dau + 1, sh.max_row + 1):
            ma = str(sh.cell(r, 1).value or '').strip()
            if not MA_DON.match(ma):
                continue
            moi = (ws_ngay.get(ma) or ship.get(ma)) if uu_tien == 'wholesale' else (ship.get(ma) or ws_ngay.get(ma))
            if moi is None:
                thieu.append(ma)
                continue
            o = sh.cell(r, cot)
            cu = '' if o.value is None else str(o.value)
            if cu[:10] == moi:
                continue
            o.value = moi
            doi.append((ma, cu, moi, '⚠ suy luận' if ma in suy_luan else ''))
        print(f'{path} / {ten_sheet}: sửa {len(doi)} ô')
        for ma, cu, moi, co in doi:
            print(f'   {ma}  {cu!r} → {moi}  {co}')
        if thieu:
            print(f'   ⚠ không có ngày nguồn cho: {", ".join(thieu)}')
        if doi:
            wb.save(path)


main()
